"""Fills the AI Builder Cup submission template with the Kakwe content.

Usage: uv run --with python-pptx --with pillow python build_pptx.py TEMPLATE.pptx OUT.pptx
Run assets.py first; it writes the images into build/.
"""
import copy
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ART = HERE / 'build'

INK = RGBColor(0x1F, 0x1A, 0x17)
INK2 = RGBColor(0x57, 0x50, 0x4A)
MUTED = RGBColor(0x6B, 0x63, 0x5C)
SAND = RGBColor(0xF7, 0xF3, 0xEE)
LINE = RGBColor(0xE6, 0xDE, 0xD3)
HIGH = RGBColor(0x9F, 0x12, 0x39)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = 'Google Sans Flex'
MY_FONT = 'A10_YoeShin'
TH_FONT = 'Noto Sans Thai'

TOP = 1.2
LEFT = 0.37
RIGHT = 9.63


def set_typeface(rpr, face):
    for tag in ('a:latin', 'a:ea', 'a:cs', 'a:sym'):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set('typeface', face)


def script_face(text):
    if any('က' <= c <= '႟' for c in text):
        return MY_FONT
    if any('฀' <= c <= '๿' for c in text):
        return TH_FONT
    return FONT


def add_runs(par, runs, size, color=INK, bold=False):
    """runs: str or list of str / (text, opts) tuples; opts may set bold, color, size, link."""
    if isinstance(runs, str):
        runs = [runs]
    for item in runs:
        text, opts = (item, {}) if isinstance(item, str) else item
        r = par.add_run()
        r.text = text
        f = r.font
        f.size = Pt(opts.get('size', size))
        f.bold = opts.get('bold', bold)
        f.color.rgb = opts.get('color', color)
        set_typeface(r._r.get_or_add_rPr(), opts.get('face', script_face(text)))
        if opts.get('link'):
            r.hyperlink.address = opts['link']


def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=4, line=1.1, inset=0.0):
    """paras: list of runs-specs, or (runs, {size, bold, color, space}) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, Inches(inset))
    for i, p in enumerate(paras):
        runs, opts = (p, {}) if not (isinstance(p, tuple) and len(p) == 2 and isinstance(p[1], dict)) else p
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = opts.get('align', align)
        par.line_spacing = opts.get('line', line)
        if i:
            par.space_before = Pt(opts.get('space', space))
        add_runs(par, runs, opts.get('size', size), opts.get('color', color), opts.get('bold', bold))
    return tb


def box(slide, x, y, w, h, fill=SAND, line=LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    style = s._element.find(qn('p:style'))
    if style is not None:
        s._element.remove(style)
    s.text_frame.text = ''
    return s


def card(slide, x, y, w, h, title, body, fill=SAND, title_size=12.5, body_size=10.5, title_color=INK, kicker=None):
    box(slide, x, y, w, h, fill=fill)
    paras = []
    if kicker:
        paras.append(([(kicker, {'bold': True, 'color': MUTED, 'size': 8.5})], {}))
    paras.append(([(title, {'bold': True, 'color': title_color, 'size': title_size})], {'space': 2 if kicker else 0}))
    paras.append(([(body, {'color': INK2, 'size': body_size})], {'space': 4}))
    text(slide, x + 0.14, y + 0.12, w - 0.28, h - 0.2, paras)


def arrow(slide, x1, y1, x2, y2, color=INK, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')
    return c


def picture(slide, name, x, y, w=None, h=None, border=False):
    path = ART / f'{name}.png'
    if w is None or h is None:
        iw, ih = Image.open(path).size
        if w is None:
            w = h * iw / ih
        else:
            h = w * ih / iw
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic, w, h


def find_shape(slide, startswith):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith(startswith):
            return sh
    raise KeyError(startswith)


def append_after(par, value, size=None, link=None, bold=None):
    """Adds a run after the paragraph's last run, copying its formatting."""
    last = par.runs[-1]._r
    r = copy.deepcopy(last)
    r.find(qn('a:t')).text = value
    last.addnext(r)
    rpr = r.find(qn('a:rPr'))
    for old in rpr.findall(qn('a:hlinkClick')):
        rpr.remove(old)
    if size:
        rpr.set('sz', str(int(size * 100)))
    if bold is not None:
        rpr.set('b', '1' if bold else '0')
    run = par.runs[-1]
    if link:
        run.hyperlink.address = link
    return run


def note(slide, s, y=5.12):
    text(slide, LEFT, y, RIGHT - LEFT, 0.3, [[(s, {'color': MUTED, 'size': 8.5})]])


def delete_slide(prs, index):
    sld = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)


# Slides


def team_details(s):
    tb = find_shape(s, 'Team Details')
    tb.width = Inches(8.0)
    tb.top = Inches(3.0)
    blank = tb.text_frame.paragraphs[1]
    if not ''.join(r.text for r in blank.runs).strip():
        blank._p.getparent().remove(blank._p)
    pars = tb.text_frame.paragraphs
    for p in pars:
        t = ''.join(r.text for r in p.runs)
        if t.startswith('Team name'):
            append_after(p, 'Super Power Rangers')
        elif t.startswith('Team leader'):
            append_after(p, 'Pyae Sone Myo (Ricky)')
            append_after(p, '   Members: Min Thein Kyaw, May Myat Noe Aung', size=12)
        elif t.startswith('Problem Statement'):
            append_after(p, 'Recruitment ads on Facebook, Viber and Telegram lure young job seekers from Myanmar and '
                            'the region into forced online scam work; OHCHR says at least 120,000 people in Myanmar may '
                            'be held in it. Kakwe helps them see the warning signs, in their own language, before they '
                            'pay, send documents or travel.', size=12)
    picture(s, 'medallion', 8.55, 3.15, w=1.05)


def brief(s):
    paras = [
        ([('Kakwe checks a job offer before you pay, send documents or travel.', {'bold': True, 'size': 15, 'color': INK})], {}),
        ([('Kakwe (', {}), ('ကာကွယ်', {}), (', Burmese for "protect") is an Android app and a web demo for migrant '
          'job seekers from Myanmar, Thailand, Indonesia and the Philippines.', {})], {'space': 10}),
        (['Paste a job ad, chat, link or screenshot, or tap the floating doll over any Android app. Gemini reads the '
          'evidence, Google Search checks named companies, and fixed rules set the verdict: Low risk, Be careful or '
          'High risk.'], {'space': 8}),
        (['The report quotes the user\'s own words, gives a verified hotline and speaks the warning in English, '
          'Burmese, Bahasa Indonesia, Tagalog or Thai.'], {'space': 8}),
        ([('The package name and web addresses keep the earlier project name saferoute.', {'color': MUTED, 'size': 9.5})],
         {'space': 12}),
    ]
    text(s, LEFT, TOP + 0.1, 5.0, 3.9, paras, size=12.5, color=INK2, line=1.15)
    picture(s, 'ph-danger-en', 5.75, TOP + 0.05, h=3.95)
    picture(s, 'ph-danger-my', 7.75, TOP + 0.05, h=3.95)


def opportunities(s):
    tb = find_shape(s, 'Opportunities')
    tf = tb.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in first.runs:
        r.text = r.text.rstrip('\x0b').rstrip()
    for br in first._p.findall(qn('a:br')):
        first._p.remove(br)
    tb.height = Inches(0.62)
    cols = [
        ('How different is it from any of the other existing ideas?',
         'It reads whole chat screenshots in Burmese and four other languages, not only links or numbers. Job Shield '
         'puts the job ad next to the recruiter\'s chat and shows each gap. Every warning quotes the words behind it.'),
        ('How will it be able to solve the problem?',
         'It works at the moment of risk: before paying, sending a passport or travelling. One tap on the floating '
         'doll checks any Android app. The report gives the next step, a verified hotline and a way to warn family.'),
        ('USP of the proposed solution',
         'Gemini finds evidence; deterministic rules decide. Wording stays in likelihood, never accusation. A Burmese '
         'voice companion answers from the finished report. The benchmark and its dataset are public.'),
    ]
    w, gap, y = 2.95, 0.18, TOP + 0.05
    for i, (q, a) in enumerate(cols):
        x = LEFT + i * (w + gap)
        box(s, x, y, w, 3.85)
        text(s, x + 0.18, y + 0.18, w - 0.36, 3.5, [
            ([(q, {'bold': True, 'size': 14})], {}),
            ([(a, {'color': INK2, 'size': 13})], {'space': 12}),
        ], line=1.15)


def features(s):
    items = [
        ('Check anything', 'Paste a message or link, or share a screenshot. Get Low risk, Be careful or High risk, with the quoted words.'),
        ('Job Shield', 'Compares the job ad with what the recruiter asks in chat and shows each gap.'),
        ('Screen bubble', 'Android. A floating doll checks whatever is on screen, in any app, with one tap.'),
        ('Scam radar', 'Shows which scams people are reporting right now, from opt-in anonymous signals.'),
        ('Voice companion', 'Talk to Pyit Taing Htaung about a report, in your language, on Gemini Live.'),
        ('Warn my family', 'Sends the warning to a trusted contact by SMS or the share sheet.'),
        ('Scam network', 'Opt-in. Repeat numbers and handles, kept as keyed hashes, raise the next warning.'),
        ('Five languages', 'English, Burmese, Bahasa Indonesia, Tagalog and Thai. Quick exit clears the screen.'),
    ]
    w, h, gx, gy = 2.17, 1.86, 0.15, 0.15
    for i, (t, b) in enumerate(items):
        x = LEFT + (i % 4) * (w + gx)
        y = TOP + 0.05 + (i // 4) * (h + gy)
        card(s, x, y, w, h, t, b, title_size=14, body_size=11.5)


def process_flow(s):
    lanes = [
        ("Ma Su's journey (a fictional story)", [
            'A Viber message offers a factory job in Thailand',
            'She pastes it into Check, or taps the floating doll',
            'High risk. Her own words are quoted as warning signs',
            'Job Shield: the ad says no fees, the chat asks for a deposit',
            'She asks Pyit Taing Htaung, in Burmese',
            'She warns her mother. She does not pay or travel',
        ], 2),
        ('The check flow, within a 15 second budget', [
            'Text, link or screenshot, with the user\'s language',
            'Gemini extraction: flags, quotes and links as JSON',
            'In parallel: Search grounding, link checks, scam network',
            'Likelihood guard rewrites certainty words',
            'Rules engine sets the score and the verdict',
            'Report, spoken warning (TTS) and voice chat (Live)',
        ], None),
    ]
    w, gap, h = 1.36, 0.22, 1.3
    y = TOP
    for title, steps, accent in lanes:
        text(s, LEFT, y, 9, 0.3, [[(title, {'bold': True, 'size': 12})]])
        by = y + 0.36
        for i, step in enumerate(steps):
            x = LEFT + i * (w + gap)
            fill = HIGH if i == accent else (INK if title.startswith('The check') and i in (1, 4) else SAND)
            fg = WHITE if fill in (HIGH, INK) else INK
            box(s, x, by, w, h, fill=fill, line=None if fill != SAND else LINE)
            text(s, x + 0.08, by + 0.06, w - 0.16, h - 0.12, [[(step, {'color': fg})]], size=11,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.08)
            if i < len(steps) - 1:
                arrow(s, x + w + 0.03, by + h / 2, x + w + gap - 0.03, by + h / 2)
        y = by + h + 0.3


def wireframes(s):
    items = [('mock-home-off', 'Home, shield off'), ('mock-home-on', 'Home, shield on'),
             ('mock-check-message', 'Check, message'), ('mock-check-job', 'Check, job offer filled')]
    h = 3.45
    widths = []
    for name, _ in items:
        iw, ih = Image.open(ART / f'{name}.png').size
        widths.append(h * iw / ih)
    gap = 0.35
    x = (10 - sum(widths) - gap * (len(items) - 1)) / 2
    for (name, cap), w in zip(items, widths):
        picture(s, name, x, TOP + 0.02, w=w, h=h, border=True)
        text(s, x - 0.1, TOP + h + 0.08, w + 0.2, 0.3, [[(cap, {'bold': True})]], size=10.5, align=PP_ALIGN.CENTER)
        x += w + gap
    note(s, 'Approved design mockups for Home (one big shield button) and the Check v3 page with kind chips, in '
            'English. The logo row and the footer line are cropped.')


def architecture(s):
    picture(s, 'arch', LEFT + 0.05, TOP + 0.02, w=9.15)
    y = 4.4
    box(s, LEFT, y, 9.26, 0.8)
    text(s, LEFT + 0.18, y + 0.1, 8.9, 0.65, [
        ([('Hosting', {'bold': True, 'size': 11})], {}),
        ([('API on Cloud Run (asia-southeast1, Singapore). Firebase for sign-in, data, web hosting and test builds.',
           {'color': INK2})], {'space': 2}),
    ], size=11)


def technologies(s):
    rows = [
        ('Gemini API, gemini-3.5-flash-lite', 'Multimodal extraction with a JSON response schema; Job Shield fact extraction'),
        ('Grounding with Google Search', 'Checks that named agencies and companies exist'),
        ('Gemini TTS, gemini-3.8-flash-tts', 'Spoken warnings in five languages'),
        ('Gemini Live API, gemini-3.8-live', 'The voice companion, Pyit Taing Htaung'),
        ('Web Risk API, Safe Browsing API', 'Known malicious links'),
        ('Firebase Auth, Firestore', 'Anonymous and Google sign-in, synced history and contacts, the scam network'),
        ('Firebase Hosting, App Distribution', 'The web demo and Android test builds'),
        ('Cloud Run, asia-southeast1', 'Hosts the check API and the voice relay'),
        ('Expo 52, React Native 0.76, TypeScript', 'The Android app and the web demo'),
        ('Hono on Node 22, TypeScript', 'The check API, the rules engine and the voice relay'),
        ('Kotlin Android module', 'The screen bubble overlay'),
    ]
    shape = s.shapes.add_table(len(rows) + 1, 2, Inches(LEFT), Inches(TOP), Inches(9.26), Inches(0.33 * (len(rows) + 1)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    for attr in ('bandRow', 'firstRow'):
        tblPr.set(attr, '0')
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'
    tbl.columns[0].width = Inches(3.4)
    tbl.columns[1].width = Inches(5.86)
    for r, (a, b) in enumerate([('Technology', 'Use in Kakwe')] + rows):
        tbl.rows[r].height = Inches(0.33)
        for c, val in enumerate((a, b)):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if r == 0 else (SAND if r % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.paragraphs[0].text = ''
            add_runs(tf.paragraphs[0], [(val, {'bold': r == 0 or c == 0, 'color': WHITE if r == 0 else (INK if c == 0 else INK2)})], 10.5)


def cost(s):
    cols = [
        ('Model calls per check', [
            'One Gemini extraction call (gemini-3.5-flash-lite), text plus an optional image.',
            'A second Gemini call with Google Search grounding, only when a company, agency or website is named.',
            'One Web Risk or Safe Browsing lookup and one RDAP lookup per link.',
            'Gemini TTS only when a spoken warning is played.',
            'A voice session lasts at most 120 seconds and counts as two checks.',
        ]),
        ('Cost controls in the API', [
            'An identical check within 10 minutes comes from the result cache.',
            'Hourly limits: 20 checks anonymous, 60 signed in, and 30 requests a minute per IP.',
            'Each check has a 15 second budget; a late stage returns "unknown".',
            'The Firebase project serves the web demo with no billing account linked.',
            'We have not measured a cost per check yet, so this deck gives no dollar figure.',
        ]),
    ]
    w = 4.55
    for i, (title, items) in enumerate(cols):
        x = LEFT + i * (w + 0.16)
        box(s, x, TOP + 0.02, w, 3.85)
        paras = [([(title, {'bold': True, 'size': 13})], {})]
        paras += [([('•  ' + t, {'color': INK2})], {'space': 7}) for t in items]
        text(s, x + 0.18, TOP + 0.17, w - 0.36, 3.6, paras, size=12, line=1.12)


def snapshots(s):
    items = [('ph-home', 'Home: the shield circle'), ('ph-check', 'Check: pick what to check'),
             ('ph-danger-en', 'High risk report'), ('ph-jobshield', "Job Shield: what doesn't match"),
             ('ph-bubble', 'Screen bubble (Android)')]
    h = 3.3
    iw, ih = Image.open(ART / 'ph-danger-en.png').size
    w = h * iw / ih
    gap = (9.26 - 5 * w) / 4
    for i, (name, cap) in enumerate(items):
        x = LEFT + i * (w + gap)
        picture(s, name, x, TOP - 0.02, w=w, h=h)
        text(s, x - 0.15, TOP + h + 0.04, w + 0.3, 0.42, [[(cap, {'bold': True})]], size=10, align=PP_ALIGN.CENTER)
    note(s, 'Web build of the release commit against the live API, with fictional inputs. The bubble shows the Android '
            'overlay views on a drawn chat.')


def benchmarks(s):
    stats = [
        ('31/31', 'Text eval cases correct', '0 responses with certainty wording after the guard. p50 2.8 s, p95 4.4 s.', True),
        ('89.7', 'Protection Score, 24 synthetic screenshots', '5 languages, 0 critical failures, 48/48 quotes found in the image.', False),
        ('about 67', 'Protection Score, 96 real screenshots', 'Provisional. Draft labels from two models; no human review yet.', False),
        ('10/10', 'Job Shield ad and chat pairs', 'No legit pair reached Pause. No scam pair stayed at the lowest tier.', False),
        ('93 + 28', 'Public dataset on Hugging Face', 'Real redacted and synthetic screenshots with labels.', False),
        ('98%', 'Verdict stability, 210 variants', 'Low resolution, JPEG, crop, rotation, dark mode, overlays and prompt injection.', False),
    ]
    w, h, gx, gy = 2.98, 1.8, 0.16, 0.16
    for i, (num, lbl, sub, dark) in enumerate(stats):
        x = LEFT + (i % 3) * (w + gx)
        y = TOP + 0.02 + (i // 3) * (h + gy)
        box(s, x, y, w, h, fill=INK if dark else SAND, line=None if dark else LINE)
        fg = WHITE if dark else INK
        text(s, x + 0.18, y + 0.12, w - 0.36, h - 0.2, [
            ([(num, {'bold': True, 'size': 28, 'color': fg})], {'line': 1.0}),
            ([(lbl, {'bold': True, 'size': 11.5, 'color': fg})], {'space': 4}),
            ([(sub, {'size': 9.5, 'color': RGBColor(0xCF, 0xC5, 0xB9) if dark else INK2})], {'space': 3}),
        ])
    note(s, 'gemini-3.5-flash-lite, production prompt, runs of 16 October 2026. A live link check also flags '
            'Google\'s Safe Browsing test phishing URL as High risk.', y=5.13)


def future(s):
    road = ['Native review of every Burmese string', 'Human review of the real-screenshot labels',
            'More languages for the region', 'Pilots with migrant worker NGOs']
    limits = ['Burmese app strings are machine translated and not yet reviewed by a native speaker.',
              'Real-screenshot labels are drafts from two models, so that score is provisional.',
              'The web demo has no screen bubble; it needs the Android app.']
    text(s, LEFT, TOP, 4.4, 0.3, [[('Next', {'bold': True, 'size': 13})]])
    for i, t in enumerate(road):
        y = TOP + 0.42 + i * 0.68
        box(s, LEFT, y, 4.45, 0.56)
        num = box(s, LEFT + 0.12, y + 0.11, 0.34, 0.34, fill=INK, line=None, shape=MSO_SHAPE.OVAL)
        text(s, LEFT + 0.12, y + 0.11, 0.34, 0.34, [[(str(i + 1), {'bold': True, 'color': WHITE})]], size=10,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, LEFT + 0.6, y + 0.04, 3.75, 0.48, [[(t, {'bold': True})]], size=11.5, anchor=MSO_ANCHOR.MIDDLE)
    x = LEFT + 4.75
    text(s, x, TOP, 4.4, 0.3, [[('Known limits today', {'bold': True, 'size': 13})]])
    box(s, x, TOP + 0.42, 4.5, 2.6)
    text(s, x + 0.18, TOP + 0.55, 4.15, 2.4, [([('•  ' + t, {'color': INK2})], {'space': 10}) for t in limits],
         size=11.5, line=1.15)


def links(s):
    tb = find_shape(s, 'Provide links')
    for p in tb.text_frame.paragraphs:
        t = ''.join(r.text for r in p.runs).strip()
        if t == 'GitHub Public Repository':
            append_after(p, ': ')
            append_after(p, 'https://github.com/minthein128/scamshield', link='https://github.com/minthein128/scamshield')
        elif t == 'Final Product Link':
            append_after(p, ': web demo ')
            append_after(p, 'https://saferoute-aibc.web.app', link='https://saferoute-aibc.web.app')
            append_after(p, ', Android APK ')
            append_after(p, 'https://saferoute-aibc.web.app/saferoute.apk', link='https://saferoute-aibc.web.app/saferoute.apk')
    y = 2.55
    box(s, LEFT, y, 9.26, 1.45)
    text(s, LEFT + 0.2, y + 0.14, 8.9, 1.3, [
        ([('Public dataset: ', {'bold': True}), ('https://huggingface.co/datasets/Rickaym/saferoute-scam-screenshots',
          {'link': 'https://huggingface.co/datasets/Rickaym/saferoute-scam-screenshots', 'color': INK2})], {}),
        ([('Team: ', {'bold': True}), ('Super Power Rangers. Pyae Sone Myo (Ricky), team leader; Min Thein Kyaw; '
          'May Myat Noe Aung', {'color': INK2})], {'space': 8}),
        ([('Contact: ', {'bold': True}), ('support@akkhayar.com', {'color': INK2, 'link': 'mailto:support@akkhayar.com'})], {'space': 8}),
    ], size=12)
    picture(s, 'medallion', 8.4, 2.8, w=0.95)


def safety(s, heading_src):
    head = copy.deepcopy(heading_src._element)
    s.shapes._spTree.append(head)
    hs = s.shapes[-1]
    hs.text_frame.paragraphs[0].runs[0].text = 'Safety and privacy'
    rules = [
        ('Likelihood words', '"Shows signs of", never "is a scam". A guard rewrites certainty phrases in all five languages.'),
        ('No accusations', 'No named person, company, number or handle is called a criminal.'),
        ('Consent first', 'Sharing is off until turned on. Overlay, screen capture, microphone, notifications and contacts are asked only when needed.'),
        ('Hashed indicators', 'Phone numbers and handles are kept as one-way hashes; signals expire after 90 days.'),
        ('Quick exit', 'One button on every report clears the screen back to Home and stops audio.'),
        ('Slim sync, 18+', 'History keeps the verdict, a score band, the title and the summary, and can be turned off.'),
    ]
    w, h, gx, gy = 3.15, 1.2, 0.15, 0.13
    for i, (t, b) in enumerate(rules):
        x = LEFT + (i % 2) * (w + gx)
        y = TOP + 0.02 + (i // 2) * (h + gy)
        card(s, x, y, w, h, t, b, title_size=12, body_size=10)
    text(s, LEFT, TOP + 3 * (h + gy) + 0.02, 6.45, 0.3, [[(
        'Images and message text are not stored or logged. Personal details are redacted before a signal is kept.',
        {'color': INK2})]], size=10)
    picture(s, 'ph-settings', 7.15, TOP - 0.02, h=3.98)


def main(src, out):
    prs = Presentation(src)
    sl = list(prs.slides)
    team_details(sl[0])
    brief(sl[2])
    opportunities(sl[3])
    features(sl[4])
    process_flow(sl[5])
    wireframes(sl[6])
    architecture(sl[7])
    technologies(sl[8])
    cost(sl[9])
    snapshots(sl[10])
    benchmarks(sl[11])
    future(sl[12])
    links(sl[13])
    safety(sl[14], find_shape(sl[2], 'Brief about the idea'))
    delete_slide(prs, 1)
    prs.save(out)
    print('wrote', out, len(prs.slides), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
